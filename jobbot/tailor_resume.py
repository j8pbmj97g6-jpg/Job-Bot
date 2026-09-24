"""Generates a tailored .docx resume per job.

Guardrail, deliberately non-negotiable: the model is instructed to only
reorder, re-emphasize, and re-word your EXISTING experience to mirror a
job's language -- never to add employers, dates, titles, skills, or
achievements that aren't in your original resume. Claiming skills you don't
have gets caught in interviews and can sink an offer after the fact, which
costs you far more than a slightly-lower ATS keyword match today.

Output: one .docx per job in tailored_resumes/, named by company + title.
The applier picks these up automatically via job['resume_path'].
"""

from __future__ import annotations

import json
import os
import re
import time
from pathlib import Path

import requests
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .db import Store

API_URL = "https://api.anthropic.com/v1/messages"
MODEL = "claude-sonnet-5"
OUT_DIR = Path("tailored_resumes")

SYSTEM_PROMPT = """You tailor a resume to a specific job posting.

STRICT RULES -- violating any of these makes the output unusable:
1. Do NOT invent or add any employer, job title, date range, degree, certification,
   or skill that is not already present in the original resume.
2. Do NOT change employment dates, company names, or job titles.
3. You MAY: reorder skills to put job-relevant ones first, rewrite bullet points
   using the job posting's terminology WITHOUT changing the underlying fact or
   outcome, rewrite the summary paragraph to foreground relevant experience,
   and drop bullets that are truly irrelevant to this job (never delete an
   entire job entry).
4. Every bullet in your output must be traceable to a real bullet in the
   original resume -- a rewording of an existing achievement, not a new one.
5. Keep it concise: at most 4 bullets per role, each under 25 words. This
   keeps output size predictable regardless of how long the job posting is.

Respond with ONLY a JSON object, no markdown fences, no other text. All
string values must be valid JSON strings (escape internal quotes as \\",
never leave a literal newline inside a string):
{
  "summary": "2-4 sentence tailored summary paragraph",
  "skills": ["ordered", "list", "of", "skills", "from", "the", "original", "resume"],
  "experience": [
    {
      "title": "exact original title",
      "company": "exact original company",
      "dates": "exact original dates",
      "location": "exact original location",
      "bullets": ["tailored bullet 1", "tailored bullet 2", "..."]
    }
  ]
}
Keep the same experience entries, in the same order, as the original resume."""


def _extract_json(text: str) -> dict:
    text = text.strip().removeprefix("```json").removeprefix("```").removesuffix("```").strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def _call_claude(resume_text: str, job: dict, api_key: str) -> dict:
    desc = (job.get("description") or "")[:6000]
    prompt = f"""ORIGINAL RESUME:
{resume_text}

JOB POSTING:
Title: {job.get('title', '')}
Company: {job.get('company', '')}
Description:
{desc}

Produce the tailored JSON now."""

    resp = requests.post(
        API_URL,
        headers={
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        json={
            "model": MODEL,
            "max_tokens": 4096,
            "system": SYSTEM_PROMPT,
            "messages": [{"role": "user", "content": prompt}],
        },
        timeout=60,
    )
    if not resp.ok:
        # The API's error body names the actual problem (bad param, invalid
        # request shape, etc.) -- swallowing it behind raise_for_status()
        # leaves only "400 Client Error", which is useless for debugging.
        raise ValueError(f"{resp.status_code} {resp.reason}: {resp.text[:500]}")
    data = resp.json()
    if data.get("stop_reason") == "max_tokens":
        raise ValueError(
            "Response was cut off (hit max_tokens) -- resume may be too long for this job. "
            "Try again or trim the job description length."
        )
    text = "".join(
        b.get("text", "") for b in data.get("content", []) if b.get("type") == "text"
    )
    return _extract_json(text)


def _safe_filename(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_")[:60]


def _render_docx(tailored: dict, contact: dict, out_path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(contact.get("name", ""))
    run.bold = True
    run.font.size = Pt(18)

    contact_line = doc.add_paragraph()
    contact_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_line.add_run(contact.get("line", ""))

    doc.add_paragraph()  # spacer

    h = doc.add_paragraph()
    h.add_run("SUMMARY").bold = True
    doc.add_paragraph(tailored.get("summary", ""))

    h = doc.add_paragraph()
    h.add_run("SKILLS").bold = True
    doc.add_paragraph(" • ".join(tailored.get("skills", [])))

    h = doc.add_paragraph()
    h.add_run("EXPERIENCE").bold = True
    for job in tailored.get("experience", []):
        p = doc.add_paragraph()
        p.add_run(f"{job.get('title', '')} — {job.get('company', '')}").bold = True
        meta = doc.add_paragraph()
        meta.add_run(f"{job.get('dates', '')}  |  {job.get('location', '')}").italic = True
        for bullet in job.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")
        doc.add_paragraph()  # spacer between roles

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def tailor_pending(
    store: Store,
    resume_text: str,
    contact: dict,
    statuses: tuple[str, ...] = ("queued",),
    limit: int | None = None,
) -> dict:
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise SystemExit("ANTHROPIC_API_KEY is not set. Run: export ANTHROPIC_API_KEY=your-key")

    # `limit` counts jobs actually tailored, not jobs merely considered --
    # fetch the full candidate pool first, filter out ones already tailored,
    # THEN cap to limit (mirrors the same fix made in cli.py's apply --limit).
    jobs = [j for s in statuses for j in store.jobs(status=s, limit=10_000)]
    jobs = [j for j in jobs if not j["resume_path"]]  # skip ones already tailored
    if limit:
        jobs = jobs[:limit]
    if not jobs:
        print("No jobs need tailoring (either none queued, or all already have a resume).")
        return {"done": 0, "errors": 0}

    done = errors = 0
    for job in jobs:
        try:
            tailored = _call_claude(resume_text, dict(job), api_key)
            # Prefix with the job id: two different postings for the same
            # company+title (a common re-post/duplicate) would otherwise
            # collide on the same filename, and the second overwrite silently
            # gives the first job's resume_path someone else's tailored content.
            fname = (f"{job['id']}_{_safe_filename(job['company'] or 'company')}"
                     f"_{_safe_filename(job['title'] or 'role')}.docx")
            out_path = OUT_DIR / fname
            _render_docx(tailored, contact, out_path)
            store.set_resume_path(job["id"], str(out_path.resolve()))
            done += 1
            print(f"  {job['id']:>5} -> {out_path}")
        except Exception as e:
            errors += 1
            print(f"  {job['id']:>5} !! tailoring failed: {e}")
        time.sleep(0.5)

    return {"done": done, "errors": errors}
